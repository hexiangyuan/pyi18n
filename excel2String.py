import openpyxl
from xml.dom.minidom import Document

items = {}


def writeToXml(values):
    doc = Document()
    resources = doc.createElement("resources")
    doc.appendChild(resources)
    for key, value in values.items():
        eString = doc.createElement("string")
        eString.setAttribute("name", key)
        print(key, "=====", value)
        if value:
            eString.appendChild(doc.createTextNode(value))
        resources.appendChild(eString)

    filename = "string-tr.xml"
    f = open(filename, "w")
    f.write(doc.toprettyxml())
    f.close()


def readExcel():
    workbook = openpyxl.load_workbook("XT APP Android TR.xlsx")
    sheet = workbook["Andriod"]
    key = "A2:C" + str(sheet.max_row)
    cell = sheet[key]
    for i in cell:
        items[i[0].value] = i[2].value


readExcel()
writeToXml(items)
